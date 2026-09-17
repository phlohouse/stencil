"""Tests for validation rule enforcement on extracted values."""

from __future__ import annotations

import openpyxl
import pytest
import yaml

from stencilpy import Stencil, ValidationError
from stencilpy.cli import main
from stencilpy.schema import FieldDef, ValidationDef
from stencilpy.validation import (
    FieldViolation,
    check_values,
    collect_violations,
    format_violations,
    is_empty,
    matches_scalar,
)


def _field(
    name: str = "value",
    *,
    type_str: str = "str",
    validation: ValidationDef | None = None,
    cell: str | None = "B2",
    range_ref: str | None = None,
) -> FieldDef:
    return FieldDef(
        name=name,
        cell=cell if range_ref is None else None,
        range=range_ref,
        type_str=type_str,
        validation=validation,
    )


class TestCollectViolations:
    def test_required_scalar_missing(self):
        field = _field(validation=ValidationDef(required=True))
        violations = collect_violations({"value": field}, {"value": None})
        assert len(violations) == 1
        assert violations[0].rule == "required"
        assert "no value" in violations[0].message

    def test_required_scalar_blank_string(self):
        field = _field(validation=ValidationDef(required=True))
        assert collect_violations({"value": field}, {"value": "   "}) != []

    def test_optional_scalar_missing_passes(self):
        field = _field(validation=ValidationDef(required=False))
        assert collect_violations({"value": field}, {"value": None}) == []

    def test_no_validation_rules_is_not_checked(self):
        field = _field(validation=None)
        assert collect_violations({"value": field}, {"value": None}) == []

    def test_min_and_max(self):
        field = _field(type_str="float", validation=ValidationDef(min=0, max=10))
        assert collect_violations({"value": field}, {"value": -1})[0].rule == "min"
        assert collect_violations({"value": field}, {"value": 11})[0].rule == "max"
        assert collect_violations({"value": field}, {"value": 10}) == []

    def test_pattern(self):
        field = _field(validation=ValidationDef(pattern=r"^[A-Z]{2}\d+$"))
        assert collect_violations({"value": field}, {"value": "AB12"}) == []
        assert collect_violations({"value": field}, {"value": "12AB"})[0].rule == "pattern"

    def test_pattern_on_numbers_uses_the_text_form(self):
        field = _field(type_str="int", validation=ValidationDef(pattern=r"^\d+$"))
        assert collect_violations({"value": field}, {"value": 1234}) == []

    def test_min_ignores_values_that_are_not_numbers(self):
        field = _field(validation=ValidationDef(min=0))
        assert collect_violations({"value": field}, {"value": "abc"}) == []

    def test_list_rules_apply_to_each_item(self):
        field = _field(
            name="readings",
            type_str="list[float]",
            cell=None,
            range_ref="D5:D",
            validation=ValidationDef(min=0, max=1000),
        )
        violations = collect_violations({"readings": field}, {"readings": [1.0, -2.0, 1500.0]})
        assert [v.field for v in violations] == ["readings[1]", "readings[2]"]
        assert [v.rule for v in violations] == ["min", "max"]

    def test_list_gaps_are_skipped(self):
        field = _field(
            name="readings",
            type_str="list[float]",
            cell=None,
            range_ref="D5:D",
            validation=ValidationDef(min=0),
        )
        assert collect_violations({"readings": field}, {"readings": [1.0, None, ""]}) == []

    def test_list_required_checks_the_field_not_the_items(self):
        field = _field(
            name="readings",
            type_str="list[float]",
            cell=None,
            range_ref="D5:D",
            validation=ValidationDef(required=True),
        )
        assert collect_violations({"readings": field}, {"readings": [None, None]}) == []
        violations = collect_violations({"readings": field}, {"readings": []})
        assert [v.rule for v in violations] == ["required"]

    def test_table_only_checks_required(self):
        field = _field(
            name="rows",
            type_str="table",
            cell=None,
            range_ref="A1:C",
            validation=ValidationDef(min=0, max=1),
        )
        assert collect_violations({"rows": field}, {"rows": [{"a": 5}]}) == []
        violations = collect_violations({"rows": field}, {"rows": []})
        assert [v.rule for v in violations] == ["required"]

    def test_computed_fields_are_checked_too(self):
        field = FieldDef(
            name="bmi",
            computed="{weight} / ({height} ** 2)",
            validation=ValidationDef(max=40),
        )
        assert collect_violations({"bmi": field}, {"bmi": 55.0})[0].rule == "max"


class TestHelpers:
    def test_is_empty(self):
        assert is_empty(None)
        assert is_empty("  ")
        assert is_empty([])
        assert is_empty({})
        assert not is_empty(0)
        assert not is_empty("x")
        assert not is_empty([None])

    def test_matches_scalar_without_rules(self):
        assert matches_scalar(_field(validation=None), "anything")

    def test_format_violations_singular_and_plural(self):
        one = [FieldViolation("a", "required", None, "no value found in the workbook")]
        assert "1 validation rule failed" in format_violations(one)
        assert format_violations(one, source="book.xlsx").startswith(
            "1 validation rule failed in 'book.xlsx'"
        )
        two = one * 2
        assert "2 validation rules failed" in format_violations(two)

    def test_check_values_raises(self):
        field = _field(validation=ValidationDef(required=True))
        with pytest.raises(ValidationError) as excinfo:
            check_values({"value": field}, {"value": None}, source="book.xlsx")
        assert "value" in str(excinfo.value)


def _write_workbook(path, **cells):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "v2.0"
    for ref, value in cells.items():
        ws[ref] = value
    wb.save(str(path))
    return path


@pytest.fixture
def bad_excel(tmp_dir):
    """A v2 workbook that breaks every kind of rule the schema declares."""
    return _write_workbook(
        tmp_dir / "lab_bad_values.xlsx",
        B3="Jane 123 Doe",  # breaks the ^[A-Za-z ]+$ pattern
        D5=1.5,
        D6=5000.0,  # breaks max: 1000
        # B10 (lab_id) is missing
    )


@pytest.fixture
def loose_excel(tmp_dir):
    """A v2 workbook that only breaks the rules the model does not enforce."""
    return _write_workbook(
        tmp_dir / "lab_loose.xlsx",
        B3="Jane Doe",
        D5=1.5,
        D6=5000.0,  # breaks max: 1000
        # B10 (lab_id) is missing
    )


@pytest.fixture
def strict_schema_yaml(tmp_dir):
    """A schema with a pattern, a per-item range rule and a required field."""
    schema = {
        "name": "lab_report",
        "discriminator": {"cells": ["A1"]},
        "versions": {
            "v2.0": {
                "fields": {
                    "patient_name": {"cell": "B3"},
                    "lab_id": {"cell": "B10"},
                    "readings": {"range": "D5:D", "type": "list[float]"},
                },
                "validation": {
                    "patient_name": {"pattern": "^[A-Za-z ]+$"},
                    "lab_id": {"required": True},
                    "readings": {"min": 0, "max": 1000},
                },
            }
        },
    }
    path = tmp_dir / "strict.stencil.yaml"
    with open(path, "w") as f:
        yaml.dump(schema, f, default_flow_style=False)
    return path


class TestStencilValidation:
    def test_extract_ignores_unenforced_rules_by_default(self, strict_schema_yaml, loose_excel):
        report = Stencil(strict_schema_yaml).extract(loose_excel)
        assert report.readings == [1.5, 5000.0]
        assert report.lab_id is None

    def test_scalar_pattern_is_still_checked_by_the_model(self, strict_schema_yaml, bad_excel):
        with pytest.raises(ValidationError) as excinfo:
            Stencil(strict_schema_yaml).extract(bad_excel)
        assert "patient_name" in str(excinfo.value)

    def test_extract_with_validate_reports_every_violation(self, strict_schema_yaml, bad_excel):
        with pytest.raises(ValidationError) as excinfo:
            Stencil(strict_schema_yaml).extract(bad_excel, validate=True)
        message = str(excinfo.value)
        assert "3 validation rules failed in 'lab_bad_values.xlsx'" in message
        assert "patient_name" in message
        assert "lab_id" in message
        assert "readings[1]" in message

    def test_validate_only_flags_the_rules_that_are_broken(self, strict_schema_yaml, loose_excel):
        with pytest.raises(ValidationError) as excinfo:
            Stencil(strict_schema_yaml).extract(loose_excel, validate=True)
        message = str(excinfo.value)
        assert "2 validation rules failed" in message
        assert "patient_name" not in message

    def test_valid_workbook_passes_with_validate(self, strict_schema_yaml, sample_excel_v2):
        report = Stencil(strict_schema_yaml).extract(sample_excel_v2, validate=True)
        assert report.patient_name == "Jane Doe"
        assert report.lab_id == "LAB-001"

    def test_batch_validate_reports_failures(self, strict_schema_yaml, bad_excel, sample_excel_v2):
        results = Stencil(strict_schema_yaml).extract(
            [bad_excel, sample_excel_v2],
            progress=False,
            concurrent=False,
            validate=True,
        )
        assert len(results.failures) == 1
        assert isinstance(results.failures[0].error, ValidationError)
        assert len(results.successes) == 1

    def test_batch_validate_with_processes_reports_failures(
        self,
        strict_schema_yaml,
        bad_excel,
        sample_excel_v2,
    ):
        results = Stencil(strict_schema_yaml).extract(
            [bad_excel, sample_excel_v2],
            progress=False,
            concurrent=True,
            validate=True,
        )
        assert len(results.failures) == 1
        assert isinstance(results.failures[0].error, ValidationError)
        assert len(results.successes) == 1


class TestCLIValidation:
    def test_extract_without_strict_ignores_unenforced_rules(
        self,
        strict_schema_yaml,
        loose_excel,
        capsys,
    ):
        assert main(["extract", str(strict_schema_yaml), str(loose_excel)]) == 0
        assert "5000" in capsys.readouterr().out

    def test_extract_strict_fails(self, strict_schema_yaml, bad_excel, capsys):
        assert main(["extract", str(strict_schema_yaml), str(bad_excel), "--strict"]) == 1
        err = capsys.readouterr().err
        assert "lab_id" in err
        assert "readings[1]" in err

    def test_extract_strict_batch_fails_only_the_bad_file(
        self,
        strict_schema_yaml,
        bad_excel,
        sample_excel_v2,
        capsys,
    ):
        code = main(
            [
                "extract",
                str(strict_schema_yaml),
                str(bad_excel.parent),
                "--include",
                "*.xlsx",
                "--strict",
                "--no-progress",
            ]
        )
        captured = capsys.readouterr()
        assert code == 1
        assert "readings[1]" in captured.err
        assert "Jane Doe" in captured.out

    def test_validate_reports_broken_rules(self, strict_schema_yaml, bad_excel, capsys):
        assert main(["validate", str(strict_schema_yaml), str(bad_excel)]) == 1
        err = capsys.readouterr().err
        assert "rule broken: patient_name" in err
        assert "rule broken: readings[1]" in err

    def test_validate_valid_file_exits_zero(self, strict_schema_yaml, sample_excel_v2, capsys):
        assert main(["validate", str(strict_schema_yaml), str(sample_excel_v2)]) == 0
        assert "rule broken" not in capsys.readouterr().err
